"""User wording that resembles template text must pass unchanged (H-058).

Run: python3 -B -m unittest deploy/tests/test_aesg_placeholders.py
"""

from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

REPOSITORY = Path(__file__).resolve().parents[2]
BRANDING = REPOSITORY / "deploy/skills/aesg-branding"
VALIDATOR = BRANDING / "scripts/validate_artifact.py"
# The recorded task card: ten requested fields, one of them END CLIENT NAME.
TASK_CARD_FIELDS = [
    "TASK REFERENCE", "PROJECT NAME", "END CLIENT NAME", "SERVICE", "SCOPE SUMMARY",
    "LOCATION", "PROJECT SIZE", "SUBMISSION DEADLINE", "AESG CONTACT", "STATUS",
]


def load_validator():
    spec = importlib.util.spec_from_file_location("aesg_validate_artifact_h058", VALIDATOR)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


validator = load_validator()


def placeholder_errors(errors: list[str]) -> list[str]:
    return [error for error in errors if error.startswith("placeholder text")]


def validate(check, path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    evidence: dict = {}
    check(path, errors, evidence)
    return placeholder_errors(errors), evidence.get("placeholderReview", [])


class WorkbookTests(unittest.TestCase):
    def generate(self, directory: Path, rows: list[dict]) -> Path:
        spec = {
            "title": "Security consultancy task card",
            "sheets": [{
                "sheet": "Task Card",
                "columns": [{"key": "field", "label": "Field", "width": 28}, {"key": "value", "label": "Value", "width": 48}],
                "rows": rows,
            }],
        }
        spec_path = directory / "spec.json"
        output = directory / "task-card.xlsx"
        spec_path.write_text(json.dumps(spec), encoding="utf-8")
        subprocess.run(
            [sys.executable, "-B", str(REPOSITORY / "deploy/skills/xlsx/scripts/create_aesg_xlsx.py"),
             "--spec", str(spec_path), "--output", str(output), "--branding-skill-dir", str(BRANDING)],
            check=True, capture_output=True, text=True,
        )
        return output

    def test_the_recorded_task_card_passes_with_every_label_unchanged(self) -> None:
        rows = [{"field": field, "value": "Not stated in the source"} for field in TASK_CARD_FIELDS]
        rows.append({"field": "NOTES", "value": "Client name as stated in the RFP"})
        with tempfile.TemporaryDirectory() as directory:
            output = self.generate(Path(directory), rows)
            result = subprocess.run([sys.executable, "-B", str(VALIDATOR), str(output)], capture_output=True, text=True)
            report = json.loads(result.stdout)
        self.assertEqual(result.returncode, 0, report["errors"])
        self.assertTrue(report["ok"])
        self.assertTrue(any("END CLIENT NAME" in warning for warning in report["warnings"]))
        self.assertTrue(any("Client name as stated in the RFP" in warning for warning in report["warnings"]))

    def test_filler_in_a_cell_still_blocks(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = self.generate(Path(directory), [
                {"field": "SCOPE SUMMARY", "value": "Lorem ipsum dolor sit amet"},
                {"field": "CLIENT", "value": "[Client Name]"},
            ])
            errors, _ = validate(validator.validate_xlsx, output)
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0].count("Task Card!"), 2)


class DocumentTests(unittest.TestCase):
    def build(self, directory: Path, paragraphs: list[str], table_label: str | None = None) -> Path:
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        if table_label:
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = table_label
            table.cell(0, 1).text = "Sample Owner LLC"
        path = directory / "document.docx"
        document.save(path)
        return path

    def test_labels_and_sentences_pass_and_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = self.build(Path(directory), [
                "END CLIENT NAME: Sample Owner LLC",
                "Use the client name as stated in the RFP.",
                "Each section heading starts on a new page.",
            ], table_label="Client Name")
            errors, review = validate(validator.validate_docx, path)
        self.assertEqual(errors, [])
        self.assertEqual(len(review), 4)

    def test_unfilled_template_slots_still_block(self) -> None:
        cases = ["Client Name", "Section Heading II", "Dear Name Another,", "Reference Number Subject: Replace Text",
                 "Lorem ipsum dolor sit amet.", "Prepared for [Client Name].", "Insert title", "xxxxxxxxxxxxx"]
        for text in cases:
            with self.subTest(text), tempfile.TemporaryDirectory() as directory:
                errors, _ = validate(validator.validate_docx, self.build(Path(directory), [text]))
                self.assertEqual(len(errors), 1, text)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is needed to make a PDF")
    def test_pdf_lines_keep_labels_and_block_filler(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, text, blocked in (("label", "END CLIENT NAME", False), ("filler", "Lorem ipsum dolor sit amet.", True)):
                source = self.build(root, [text])
                pdf_dir = root / name
                subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir), str(source)],
                               check=True, capture_output=True)
                errors, _ = validate(validator.validate_pdf, pdf_dir / "document.pdf")
                self.assertEqual(bool(errors), blocked, text)


class PresentationTests(unittest.TestCase):
    def build(self, directory: Path, text: str) -> Path:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = text
        path = directory / "deck.pptx"
        presentation.save(path)
        return path

    def test_slide_wording_follows_the_same_rules(self) -> None:
        for text, blocked in (("END CLIENT NAME", False), ("Name Surname", True), ("Click to edit Master title style", True)):
            with self.subTest(text), tempfile.TemporaryDirectory() as directory:
                errors, _ = validate(validator.validate_pptx, self.build(Path(directory), text))
                self.assertEqual(bool(errors), blocked, text)


if __name__ == "__main__":
    unittest.main()
