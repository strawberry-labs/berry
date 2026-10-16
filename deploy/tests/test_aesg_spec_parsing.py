"""A repeated JSON key must stop generation instead of dropping content (H-056).

Run: python3 -B -m unittest deploy/tests/test_aesg_spec_parsing.py
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

REPOSITORY = Path(__file__).resolve().parents[2]
BRANDING = REPOSITORY / "deploy/skills/aesg-branding"
GENERATORS = {
    "docx": REPOSITORY / "deploy/skills/docx/scripts/create_aesg_docx.py",
    "pptx": REPOSITORY / "deploy/skills/pptx/scripts/create_aesg_pptx.py",
    "xlsx": REPOSITORY / "deploy/skills/xlsx/scripts/create_aesg_xlsx.py",
}
# The shape of the recorded failure: one section repeats "paragraphs", so the
# reference, the acceptance statement and its lead-in were replaced by the
# closing lines while the document still looked complete.
OPENING = [
    "Reference: proposal P-1234 for the sample retail project.",
    "AESG confirms that it accepts the scope described in the documents below.",
    "The accepted documents are:",
]
CLOSING = [
    "Any change to this scope will be agreed in writing.",
    "Signed for and on behalf of AESG.",
]


def run(kind: str, spec_text: str, directory: Path) -> tuple[subprocess.CompletedProcess[str], Path]:
    spec_path = directory / "spec.json"
    output = directory / f"out.{kind}"
    spec_path.write_text(spec_text, encoding="utf-8")
    command = [
        sys.executable, "-B", str(GENERATORS[kind]), "--spec", str(spec_path),
        "--output", str(output), "--branding-skill-dir", str(BRANDING),
    ]
    if kind == "docx":
        command.append("--no-toc-pages")
    return subprocess.run(command, capture_output=True, text=True), output


def section(*paragraph_lists: list[str]) -> str:
    keys = ", ".join(f'"paragraphs": {json.dumps(paragraphs)}' for paragraphs in paragraph_lists)
    return '{"heading": "Scope acceptance", ' + keys + "}"


def report(*sections: str) -> str:
    return '{"kind": "report", "title": "Scope acceptance statement", "sections": [' + ", ".join(sections) + "]}"


def docx_paragraphs(path: Path) -> list[str]:
    return [paragraph.text for paragraph in Document(path).paragraphs]


class RepeatedKeyTests(unittest.TestCase):
    def test_a_repeated_paragraphs_key_stops_before_any_output(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result, output = run("docx", report(section(OPENING, CLOSING)), Path(directory))
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('"paragraphs" in spec.sections[0]', result.stderr)
            self.assertIn("nothing was written", result.stderr)
            self.assertFalse(output.exists())

    def test_the_corrected_spec_keeps_every_paragraph(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result, output = run("docx", report(section(OPENING + CLOSING)), Path(directory))
            self.assertEqual(result.returncode, 0, result.stderr)
            text = docx_paragraphs(output)
        for paragraph in OPENING + CLOSING:
            self.assertIn(paragraph, text)

    def test_equal_values_and_nested_repeats_are_refused_with_their_path(self) -> None:
        nested = (
            '{"heading": "Actions", "blocks": [{"type": "paragraph", "text": "One."},'
            ' {"type": "paragraph", "text": "Two.", "text": "Two."}]}'
        )
        with tempfile.TemporaryDirectory() as directory:
            result, output = run("docx", report(section(OPENING), nested), Path(directory))
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('"text" in spec.sections[1].blocks[1]', result.stderr)
            self.assertFalse(output.exists())

    def test_the_same_key_in_separate_objects_is_valid(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result, output = run("docx", report(section(OPENING), section(CLOSING)), Path(directory))
            self.assertEqual(result.returncode, 0, result.stderr)
            text = docx_paragraphs(output)
        for paragraph in OPENING + CLOSING:
            self.assertIn(paragraph, text)

    def test_presentation_and_workbook_specs_are_checked_too(self) -> None:
        specs = {
            "pptx": ('{"title": "Deck", "slides": [{"layout": "text", "title": "One", "title": "Two"}]}', "spec.slides[0]"),
            "xlsx": ('{"sheets": [{"title": "Register", "columns": ["A"], "columns": ["B"], "rows": [["x"]]}]}', "spec.sheets[0]"),
        }
        for kind, (text, where) in specs.items():
            with self.subTest(kind), tempfile.TemporaryDirectory() as directory:
                result, output = run(kind, text, Path(directory))
                self.assertNotEqual(result.returncode, 0)
                self.assertIn(where, result.stderr)
                self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
